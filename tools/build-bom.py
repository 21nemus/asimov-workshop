"""Extracts the supplied Asimov 1 BOM workbook into src/data/bom.json.

Every record keeps its source identity (document, revision, sheet, row), the raw cell
text exactly as written, and the few normalised fields the app needs. Nothing is
inferred: a missing quantity stays null, a cached formula total over an empty range is
reported as unsupported, and rows that share a part number but disagree are flagged,
never merged. The workbook itself is not redistributed.

Usage: python3 -I tools/build-bom.py <workbook.xlsx> <out.json>
"""
import hashlib
import json
import re
import sys
import warnings

import openpyxl

warnings.filterwarnings('ignore')

src, out = sys.argv[1], sys.argv[2]
data = open(src, 'rb').read()
sha = hashlib.sha256(data).hexdigest()
wb_f = openpyxl.load_workbook(src, data_only=False)
wb_v = openpyxl.load_workbook(src, data_only=True)

cover = {r[0].value: r[1].value for r in wb_f['Cover'].iter_rows(min_row=2, max_row=5) if r[0].value}
DOC = cover.get('Document No.')
REV = cover.get('Revision')
SHEETS = [
    ('Manufacture - Mechanical', 'mechanical'),
    ('Manufacture - Electrical', 'electrical'),
    ('Manufacture - Motor Wiring Uniq', 'wiring'),
    ('Procure - Generic Unique', 'generic'),
    ('Procure - Actuators Unique', 'actuator'),
    ('Procure - Fasteners Unique', 'fastener'),
    ('Procure - Misc Unique', 'misc'),
]
CANCEL_MARKS = ['（已取消）', '(已取消)', '已取消']


def text(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def normalise_part(raw):
    """Search alias only: the raw value is always kept alongside."""
    if raw is None:
        return None
    s = str(raw)
    for m in CANCEL_MARKS:
        s = s.replace(m, '')
    s = s.strip().rstrip('.').strip()
    s = re.sub(r'\s+', '', s)
    return s.upper() or None


records = []
for sheet, kind in SHEETS:
    ws = wb_f[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = [text(h).strip() if h is not None and text(h).strip() else None for h in rows[0]]
    for i, row in enumerate(rows[1:], start=2):
        cells = {}
        for j, v in enumerate(row):
            if v is None or (isinstance(v, str) and not v.strip()):
                continue
            col = openpyxl.utils.get_column_letter(j + 1)
            name = header[j] if j < len(header) and header[j] else col
            cells[name] = text(v)
        if not cells:
            continue
        part_raw = cells.get('Part Name')
        qty_raw = cells.get('Qty')
        qty = None
        if qty_raw is not None:
            try:
                qty = float(qty_raw)
                qty = int(qty) if qty.is_integer() else qty
            except ValueError:
                qty = None
        flags = []
        if part_raw and any(m in part_raw for m in CANCEL_MARKS):
            flags.append('cancelled')
        if kind in ('mechanical', 'generic', 'actuator', 'fastener', 'misc') and qty_raw is None:
            flags.append('missing-quantity')
        if cells.get('Company', '').strip() == 'Company':
            flags.append('vendor-placeholder')
        records.append({
            'id': f'{DOC}/{REV}/{sheet}/row{i}',
            'sheet': sheet,
            'row': i,
            'kind': kind,
            'subAssembly': cells.get('Sub-Assembly'),
            'partNameRaw': part_raw,
            'partNumber': normalise_part(part_raw),
            'qty': qty,
            'qtyRaw': qty_raw,
            'material': cells.get('Material'),
            'finish': cells.get('Surface Finish'),
            'description': cells.get('Description'),
            'sku': cells.get('SKU'),
            'vendor': cells.get('Company'),
            'url': cells.get('Product URL'),
            'raw': cells,
            'flags': flags,
        })

# Same part number in a different module than its sub-assembly says.
for r in records:
    pn, sa = r['partNumber'], r['subAssembly']
    if pn and sa and pn.startswith('ASV1_') and not pn.startswith(sa.strip().upper() + '_'):
        r['flags'].append('part-number-module-mismatch')

# Conflicts: one part number, several records that disagree on what the part is.
by_pn = {}
for r in records:
    if r['partNumber'] and r['partNumber'].startswith('ASV1_'):
        by_pn.setdefault(r['partNumber'], []).append(r)
conflicts = []
for pn, rs in sorted(by_pn.items()):
    kinds = sorted({r['kind'] for r in rs})
    descs = sorted({(r['description'] or r['material'] or '').strip() for r in rs})
    if len(rs) > 1 and (len(kinds) > 1 or len(descs) > 1):
        conflicts.append({'partNumber': pn, 'records': [r['id'] for r in rs], 'kinds': kinds, 'descriptions': descs})
        for r in rs:
            r['flags'].append('conflicting-part-number')
    elif len(rs) > 1:
        conflicts.append({'partNumber': pn, 'records': [r['id'] for r in rs], 'kinds': kinds, 'descriptions': descs, 'note': 'repeated, same description'})

# Totals sheets: cached formula values are reported, never treated as verified totals.
totals = []
for sheet in ['Manufacture - Motor Wiring Tota', 'Procure - Generic Total', 'Procure - Actuators Total', 'Procure - Fasteners Total']:
    ws_f, ws_v = wb_f[sheet], wb_v[sheet]
    for row in ws_f.iter_rows():
        for c in row:
            if c.value in (None, ''):
                continue
            entry = {'sheet': sheet, 'cell': c.coordinate, 'value': text(c.value), 'cached': text(ws_v[c.coordinate].value)}
            m = re.match(r'^=SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)$', str(c.value))
            if m:
                col, r0, _, r1 = m.group(1), int(m.group(2)), m.group(3), int(m.group(4))
                support = [ws_f[f'{col}{r}'].value for r in range(r0, r1 + 1)]
                empty = all(v in (None, '') for v in support)
                entry['status'] = 'unsupported: the summed range is empty in the supplied workbook' if empty else 'formula over populated range (not independently verified)'
            else:
                entry['status'] = 'stray cell on an otherwise empty totals sheet'
            totals.append(entry)
    entry_pivots = len(getattr(ws_f, '_pivots', []))
    if entry_pivots:
        totals.append({'sheet': sheet, 'cell': None, 'value': None, 'cached': None, 'status': f'{entry_pivots} pivot table(s) whose cached results are not visible as cell values'})

out_doc = {
    'source': {
        'document': DOC, 'revision': REV, 'scope': cover.get('Scope'), 'project': cover.get('Project'),
        'file': 'Asimov 1 - Bill of Materials.xlsx', 'sha256': sha, 'bytes': len(data),
        'licence': 'Not stated in the workbook. Distributed by Menlo Research through an email request form. Extracted facts only; the workbook is not redistributed.',
    },
    'rules': [
        'Record identity is document/revision/sheet/row. Raw cell text is kept verbatim in `raw`.',
        '`partNumber` is a normalised search alias (cancellation marks, whitespace and trailing dots removed); never use it to merge records.',
        'Missing quantities stay null. No prices exist in the source; none are computed.',
        'Cached totals are listed with their status; none is presented as a verified component count.',
    ],
    'records': records,
    'conflicts': conflicts,
    'totals': totals,
}
json.dump(out_doc, open(out, 'w'), ensure_ascii=False, indent=1)
print(f'{len(records)} records, {len(conflicts)} repeated/conflicting part numbers, {len(totals)} totals cells')
for c in conflicts:
    print('  ', c['partNumber'], c['kinds'], c['descriptions'], c.get('note', ''))
